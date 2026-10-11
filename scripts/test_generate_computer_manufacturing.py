"""Independent export checks for the synthetic manufacturing data generator."""

import csv
import argparse
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from datetime import datetime
from collections import defaultdict
from decimal import Decimal as D
from pathlib import Path

from openpyxl import load_workbook

from generate_computer_manufacturing import Generator, IMPORT_TABLES, parser


class ManufacturingDataTests(unittest.TestCase):
    """Check financial conservation, lineage, mutations, depth and reproducibility."""

    def generate(self, output, *flags):
        args = parser().parse_args(["--output", str(output), "--batches", "1", *flags])
        Generator(args).run().write(output)
        tables = {}
        for path in output.glob("*.csv"):
            with path.open(encoding="utf-8", newline="") as stream:
                tables[path.stem] = list(csv.DictReader(stream))
        return tables

    def check_ledgers(self, tables):
        bom_keys = [(r["bom_no"], r["product_no"], r["material_no"]) for r in tables["bill_of_material"]]
        self.assertEqual(len(bom_keys), len(set(bom_keys)), "Standard BOM violates warehouse uniqueness")
        issue_keys = [(r["date"], r["movement_no"], r["material_no"]) for r in tables["inventory_usage"]]
        self.assertEqual(len(issue_keys), len(set(issue_keys)), "Actual issues violate warehouse uniqueness")
        outputs = {r["order_no"]: r for r in tables["production_outputs"]}
        self.assertEqual(len(outputs), len(tables["production_outputs"]))
        amounts, allocations, quantities, shares = (defaultdict(lambda: D(0)) for _ in range(4))
        purchases = {r["purchase_order_no"]: r for r in tables["purchases"]}
        events = {r["event_id"] for r in tables["anomalies"]}
        self.assertEqual(len(tables["inventory_usage"]), len(tables["issue_evidence"]))
        seen_events = set()
        for issue, evidence in zip(tables["inventory_usage"], tables["issue_evidence"]):
            for key in ("date", "movement_no", "order_no", "product_no", "material_no"):
                self.assertEqual(issue[key], evidence[key])
            output = outputs[issue["order_no"]]
            self.assertEqual(issue["product_no"], output["product_no"])
            self.assertEqual(issue["product_num"], output["product_num"])
            self.assertEqual(issue["product_total_cost"], output["product_total_cost"])
            self.assertGreater(D(issue["material_num"]), 0)
            amounts[issue["order_no"]] += D(issue["material_total_cost"])
            source = evidence["source_order_no"]
            if source:
                source_output = outputs[source]
                self.assertEqual(source_output["product_no"], issue["material_no"])
                self.assertLessEqual(source_output["date"], issue["date"])
                quantities[source] += D(issue["material_num"])
                shares[source] += D(evidence["source_share"])
            else:
                po = purchases[evidence["purchase_order_no"]]
                self.assertEqual(po["product_no"], issue["material_no"])
                self.assertEqual(po["product_num"], issue["material_num"])
                self.assertEqual(po["product_total_price"], issue["material_total_cost"])
                self.assertLessEqual(po["date"], issue["date"])
            seen_events.update(filter(None, evidence["event_ids"].split(";")))
        self.assertEqual(events, seen_events)
        allocated_pools = defaultdict(lambda: [D(0), D(0)])
        for row in tables["processing_allocations"]:
            allocations[row["order_no"]] += D(row["labour"]) + D(row["overhead"])
            allocated_pools[row["source_id"]][0] += D(row["labour"])
            allocated_pools[row["source_id"]][1] += D(row["overhead"])
        for pool in tables["processing_sources"]:
            self.assertEqual(allocated_pools[pool["source_id"]], [D(pool["labour"]), D(pool["overhead"])])
        for order, output in outputs.items():
            self.assertEqual(D(output["product_total_cost"]), amounts[order] + allocations[order])
        parent_amounts = defaultdict(lambda: D(0))
        for issue, evidence in zip(tables["inventory_usage"], tables["issue_evidence"]):
            if evidence["source_order_no"]:
                parent_amounts[evidence["source_order_no"]] += D(issue["material_total_cost"])
        for source in quantities:
            self.assertEqual(quantities[source], D(outputs[source]["product_num"]))
            self.assertEqual(shares[source], D(1))
            self.assertEqual(parent_amounts[source], D(outputs[source]["product_total_cost"]))
        # Recover tree depths independently from explicit source-order links.
        children = defaultdict(set)
        for row in tables["issue_evidence"]:
            children[row["order_no"]].add(row["source_order_no"] or row["position"])

        def depths(order, visiting):
            self.assertNotIn(order, visiting, "Production lineage contains a cycle")
            if order not in outputs:
                return [1]
            return [1 + depth for child in children[order]
                    for depth in depths(child, visiting | {order})]

        for expected in tables["batch_expectations"]:
            leaves = depths(expected["root_order_no"], set())
            self.assertEqual(max(leaves), int(expected["maximum_depth"]))
            self.assertLessEqual(max(leaves), 10)
            self.assertGreater(max(leaves), min(leaves))
            self.assertEqual(D(expected["actual_total_cost"]),
                             D(outputs[expected["root_order_no"]]["product_total_cost"]))
            self.assertEqual(D(expected["difference_from_nominal"]),
                             D(expected["actual_total_cost"]) - D(expected["nominal_total_cost"]))
        for sale in tables["sales_details"]:
            output = outputs[sale["movement_no"].removeprefix("SALE-")]
            self.assertEqual(sale["total_cost"], output["product_total_cost"])
        for rows in tables.values():
            self.assertTrue(all(str(value).isascii() for row in rows for value in row.values()))

    def test_default_ledgers_and_mutations(self):
        with tempfile.TemporaryDirectory() as tmp:
            tables = self.generate(Path(tmp))
            self.check_ledgers(tables)
            kinds = {r["type"] for r in tables["anomalies"]}
            self.assertEqual(kinds, {"STRUCTURE_INSERT", "STRUCTURE_EXTRA_COMPONENT",
                                    "MATERIAL_SUBSTITUTION", "LEAF_COST_CHANGE"})
            self.assertTrue(all(r["period"] == "ACTUAL" for r in tables["anomalies"]))

    def test_zero_rates_leave_costs_and_depth_unchanged(self):
        with tempfile.TemporaryDirectory() as tmp:
            tables = self.generate(Path(tmp), "--structure-rate", "0", "--material-rate", "0", "--cost-rate", "0")
            self.check_ledgers(tables)
            self.assertEqual(tables["anomalies"], [])
            self.assertTrue(all(D(r["difference_from_nominal"]) == 0 for r in tables["batch_expectations"]))
            self.assertTrue(all(r["maximum_depth"] == "10" for r in tables["batch_expectations"]))

    def test_default_four_batches_preserve_repeated_material_positions(self):
        with tempfile.TemporaryDirectory() as tmp:
            tables = self.generate(Path(tmp), "--batches", "4")
            self.check_ledgers(tables)
            issues = [row for row in tables["inventory_usage"]
                      if row["order_no"] == "MO-0004158" and row["material_no"] == "RAW-006"]
            self.assertEqual(len(issues), 4)
            self.assertEqual({row["movement_no"] for row in issues},
                             {f"ISS-MO-0004158-{split}-P{position}"
                              for split in (1, 2) for position in (1, 3)})
            self.assertEqual(sum(D(row["material_num"]) for row in issues), D("40"))
            self.assertEqual(sum(D(row["material_total_cost"]) for row in issues), D("2"))
            self.assertTrue(all(D(row["product_total_cost"]) == D("38") for row in issues))
            evidence = [row for row in tables["issue_evidence"]
                        if row["order_no"] == "MO-0004158" and row["material_no"] == "RAW-006"]
            self.assertEqual(len({row["position"] for row in evidence}), 2)
            self.assertEqual(len({row["purchase_order_no"] for row in evidence}), 4)
            self.assertEqual(len({row["allocation_id"] for row in evidence}), 2)

    def test_every_position_can_mutate_in_both_periods(self):
        with tempfile.TemporaryDirectory() as tmp:
            tables = self.generate(Path(tmp), "--structure-rate", "1", "--material-rate", "1",
                                   "--cost-rate", "1", "--baseline-scale", "1", "--max-depth", "3")
            self.check_ledgers(tables)
            self.assertEqual({r["period"] for r in tables["anomalies"]}, {"COMPARABLE", "ACTUAL"})
            recipes = defaultdict(list)
            for row in tables["bill_of_material"]:
                recipes[row["product_no"]].append(row["material_no"])

            def counts(code):
                nested = [counts(child) for child in recipes[code]]
                return (1 + sum(n[0] for n in nested), sum(n[1] for n in nested)) if nested else (1, 1)

            for expected in tables["batch_expectations"]:
                self.assertGreater(int(expected["event_count"]), 0)
                self.assertLessEqual(int(expected["maximum_depth"]), 3)
                total_nodes, leaves = counts(expected["product_no"])
                events = [r for r in tables["anomalies"] if r["root_order_no"] == expected["root_order_no"]]
                self.assertEqual(sum(r["type"] == "MATERIAL_SUBSTITUTION" for r in events), total_nodes - 1)
                self.assertEqual(sum(r["type"].startswith("STRUCTURE_") for r in events), total_nodes - 1)
                self.assertEqual(sum(r["type"] == "LEAF_COST_CHANGE" for r in events), leaves)

    def test_seed_reproduces_every_export(self):
        with tempfile.TemporaryDirectory() as tmp:
            left, right = Path(tmp) / "left", Path(tmp) / "right"
            self.generate(left)
            self.generate(right)
            self.assertEqual({p.name: p.read_bytes() for p in left.iterdir() if p.is_file()},
                             {p.name: p.read_bytes() for p in right.iterdir() if p.is_file()})
            self.assertEqual((left / "excel/input.json").read_bytes(), (right / "excel/input.json").read_bytes())
            for name in IMPORT_TABLES:
                with zipfile.ZipFile(left / "excel" / f"{name}.xlsx") as a, \
                     zipfile.ZipFile(right / "excel" / f"{name}.xlsx") as b:
                    self.assertEqual(a.namelist(), b.namelist())
                    for entry in a.namelist():
                        if entry != "docProps/core.xml":
                            self.assertEqual(a.read(entry), b.read(entry), entry)

    def test_excel_files_match_exact_csv_fixtures(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            tables = self.generate(output, "--batches", "4")
            fixture = json.loads((output / "excel/input.json").read_text(encoding="utf-8"))
            self.assertEqual(set(fixture), set(IMPORT_TABLES))
            for name in IMPORT_TABLES:
                path = output / "excel" / f"{name}.xlsx"
                book = load_workbook(path)
                try:
                    self.assertEqual(book.sheetnames, [name])
                    sheet = book[name]
                    self.assertEqual(sheet.freeze_panes, "A2")
                    self.assertFalse(sheet.sheet_view.showGridLines)
                    headers = fixture[name]["headers"]
                    self.assertEqual([c.value for c in sheet[1]], headers)
                    self.assertEqual(sheet.max_row, len(tables[name]) + 1)
                    self.assertEqual(fixture[name]["rows"], [[row[key] for key in headers] for row in tables[name]])
                    for cells, source in zip(sheet.iter_rows(min_row=2), tables[name]):
                        for cell, key in zip(cells, headers):
                            if key == "date":
                                self.assertIsInstance(cell.value, datetime)
                                self.assertEqual(cell.value.date().isoformat(), source[key])
                                self.assertEqual(cell.number_format, "yyyy-mm-dd")
                            elif cell.data_type == "n" and cell.value is not None:
                                self.assertEqual(D(str(cell.value)).quantize(D("0.000001")), D(source[key]))
                                self.assertEqual(cell.number_format, "0.000000")
                            else:
                                self.assertEqual(cell.value or "", source[key])
                    self.assertGreaterEqual(sheet.column_dimensions["A"].width, 18)
                finally:
                    book.close()
                with zipfile.ZipFile(path) as archive:
                    self.assertNotIn(b't="d"', archive.read("xl/worksheets/sheet1.xml"))

    def test_excel_preserves_six_places_and_rejects_precision_loss(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            tables = {name: [dict(unit_cost=D("1.234567"))] for name in IMPORT_TABLES}
            Generator.write_excel(output, tables)
            book = load_workbook(output / "material_master.xlsx")
            try:
                self.assertEqual(D(str(book.active["A2"].value)), D("1.234567"))
            finally:
                book.close()
            tables["material_master"][0]["unit_cost"] = D("123456789012.123456")
            with self.assertRaisesRegex(ValueError, "Excel numeric precision exceeds"):
                Generator.write_excel(output / "rejected", tables)

    def test_cli_rejects_invalid_input_and_preserves_existing_files(self):
        script = Path(__file__).with_name("generate_computer_manufacturing.py")
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "existing"
            output.mkdir()
            marker = output / "keep.txt"
            marker.write_text("Preserve this file", encoding="utf-8")
            for flags in ([], ["--material-rate", "nan"], ["--structure-rate", "-1"],
                          ["--cost-rate", "1.1"], ["--max-depth", "2"], ["--batches", "0"]):
                result = subprocess.run([sys.executable, str(script), "--output", str(output), *flags],
                                        capture_output=True, text=True)
                self.assertEqual(result.returncode, 2)
            self.assertEqual(marker.read_text(encoding="utf-8"), "Preserve this file")


if __name__ == "__main__":
    if "--dataset" in sys.argv:
        cli = argparse.ArgumentParser(description="Validate an existing generated data package")
        cli.add_argument("--dataset", type=Path, required=True)
        args = cli.parse_args()
        tables = {}
        for path in args.dataset.glob("*.csv"):
            with path.open(encoding="utf-8", newline="") as stream:
                tables[path.stem] = list(csv.DictReader(stream))
        ManufacturingDataTests().check_ledgers(tables)
        report = dict(status="PASS", checks=["Warehouse BOM and issue key uniqueness", "Decimal cost conservation", "Source quantity and share conservation",
                      "Explicit lineage and acyclic variable-depth trees", "Issue/output/purchase/sales consistency",
                      "Processing pool allocations", "Anomaly evidence references", "English ASCII values"],
                      rows={name: len(rows) for name, rows in tables.items()})
        (args.dataset / "validation.json").write_text(json.dumps(report, indent=2) + "\n", encoding="utf-8")
        print(json.dumps(report, indent=2))
    else:
        unittest.main()
